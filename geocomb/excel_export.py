from __future__ import annotations

import os
import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .constants import CLASS_ORDER
from .database import ProjectDatabase, utc_now


HEADER_FILL = PatternFill(
    fill_type="solid",
    fgColor="1F4E78",
)
HEADER_FONT = Font(
    color="FFFFFF",
    bold=True,
)


def style_sheet(sheet) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    for column in sheet.columns:
        values = [
            "" if cell.value is None else str(cell.value)
            for cell in column
        ]
        width = min(
            max(len(value) for value in values) + 2,
            45,
        )
        sheet.column_dimensions[
            get_column_letter(column[0].column)
        ].width = max(width, 10)


def build_workbook(database: ProjectDatabase) -> Workbook:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Image_Summary"

    summary.append(
        [
            "Image",
            "Total",
            *CLASS_ORDER,
            "Edited cells",
            "Added cells",
            "Deleted cells",
        ]
    )

    images = database.list_images()

    for image in images:
        image_id = int(image["id"])
        counts = database.class_counts(image_id)
        cells = database.cells_for_image(
            image_id,
            include_deleted=True,
        )
        active = [row for row in cells if not row["deleted"]]
        summary.append(
            [
                image["original_name"],
                len(active),
                *[counts[name] for name in CLASS_ORDER],
                sum(int(row["edited"]) for row in cells),
                sum(int(row["added"]) for row in cells),
                sum(int(row["deleted"]) for row in cells),
            ]
        )

    details = workbook.create_sheet("Cell_Details")
    details.append(
        [
            "Image",
            "Cell ID",
            "X (5000-px canvas)",
            "Y (5000-px canvas)",
            "X (original image)",
            "Y (original image)",
            "Estimated cell pitch (canvas px)",
            "Original class",
            "Current class",
            "Confidence",
            "Edited",
            "Added",
            "Deleted",
            "Created at",
            "Updated at",
        ]
    )

    for row in database.all_cells():
        details.append(
            [
                row["original_name"],
                row["id"],
                row["x_canvas"],
                row["y_canvas"],
                row["x_original"],
                row["y_original"],
                row["estimated_step_canvas"],
                row["class_initial"],
                row["class_current"],
                row["confidence"],
                bool(row["edited"]),
                bool(row["added"]),
                bool(row["deleted"]),
                row["created_at"],
                row["updated_at"],
            ]
        )

    history = workbook.create_sheet("Edit_History")
    history.append(
        [
            "Timestamp",
            "Image",
            "Cell ID",
            "Action",
            "Previous value",
            "New value",
        ]
    )

    for row in database.edit_history():
        history.append(
            [
                row["timestamp"],
                row["original_name"],
                row["cell_id"],
                row["action"],
                row["old_value"],
                row["new_value"],
            ]
        )

    info = workbook.create_sheet("Analysis_Info")
    info.append(["Field", "Value"])
    info_rows = [
        ("Software", "GeoComb"),
        ("Export timestamp (UTC)", utc_now()),
        (
            "Project folder",
            str(database.project_folder),
        ),
        (
            "Coordinate systems",
            (
                "Canvas coordinates refer to the 5,000-pixel "
                "masked frame crop. Original coordinates are "
                "mapped back using the detected frame bounding "
                "box and image-specific scale."
            ),
        ),
        ("Classes", ", ".join(CLASS_ORDER)),
    ]

    for key, value in info_rows:
        info.append([key, value])

    for sheet in workbook.worksheets:
        style_sheet(sheet)

    return workbook


def export_excel(
    database: ProjectDatabase,
    output_path: Path | None = None,
) -> tuple[Path, bool, str]:
    if output_path is None:
        output_path = (
            database.project_folder
            / "results"
            / "GeoComb_results.xlsx"
        )

    output_path = Path(output_path)
    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    workbook = build_workbook(database)

    with tempfile.NamedTemporaryFile(
        suffix=".xlsx",
        delete=False,
        dir=output_path.parent,
    ) as temporary:
        temporary_path = Path(temporary.name)

    try:
        workbook.save(temporary_path)

        try:
            os.replace(temporary_path, output_path)
            return output_path, True, "Excel synchronized"
        except PermissionError:
            pending = output_path.with_name(
                output_path.stem
                + "_pending_update"
                + output_path.suffix
            )
            os.replace(temporary_path, pending)
            return (
                pending,
                False,
                (
                    "The main Excel file appears to be open. "
                    "A pending-update workbook was written instead."
                ),
            )
    finally:
        if temporary_path.exists():
            temporary_path.unlink(missing_ok=True)
